"""
Report generation utilities (PDF and Excel)
"""
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
import io


def generar_pdf_nomina(nomina, empleado):
    """
    Generate PDF report for a single payroll record
    
    Args:
        nomina: Nomina object
        empleado: Empleado object
    
    Returns:
        BytesIO buffer with PDF content
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()
    
    # Title style
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#2c3e50'),
        alignment=TA_CENTER,
        spaceAfter=30
    )
    
    # Add title
    title = Paragraph("RECIBO DE NÓMINA", title_style)
    elements.append(title)
    elements.append(Spacer(1, 0.2*inch))
    
    # Employee info
    info_data = [
        ['Empleado:', empleado.nombre],
        ['Periodo:', nomina.periodo],
        ['Fecha de Cálculo:', nomina.fecha_calculo.strftime('%d/%m/%Y')],
        ['Días Trabajados:', str(nomina.dias_trabajados)]
    ]
    
    info_table = Table(info_data, colWidths=[2*inch, 4*inch])
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#ecf0f1')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#2c3e50')),
        ('ALIGN', (0, 0), (0, -1), 'RIGHT'),
        ('ALIGN', (1, 0), (1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
    ]))
    
    elements.append(info_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Payroll details
    detail_data = [
        ['CONCEPTO', 'CANTIDAD', 'MONTO'],
        ['Sueldo Base', f'{nomina.dias_trabajados} días', f'${nomina.sueldo_base:,.2f}'],
        ['Horas Extra', f'{nomina.horas_extra} hrs', f'${nomina.pago_extras:,.2f}'],
        ['Ajuste por Faltas', f'{nomina.faltas} faltas', f'${nomina.ajuste_faltas:,.2f}'],
        ['Bono de Puntualidad', '', f'${nomina.bono_puntualidad:,.2f}'],
        ['Bono de Producción', '', f'${nomina.bono_produccion:,.2f}'],
        ['', '', ''],
        ['TOTAL PERCEPCIONES', '', f'${nomina.sueldo_base + nomina.pago_extras + nomina.ajuste_faltas + nomina.total_bonos:,.2f}'],
        ['', '', ''],
        ['Tardes', f'{nomina.tardes} tardes', f'-${nomina.descuento_tardes:,.2f}'],
        ['Préstamo (3%)', f'${nomina.prestamo:,.2f}', f'-${nomina.descuento_prestamo:,.2f}'],
        ['Caja de Ahorro', '', f'-${nomina.caja_ahorro:,.2f}'],
        ['', '', ''],
        ['TOTAL DEDUCCIONES', '', f'-${nomina.total_descuentos:,.2f}'],
        ['', '', ''],
        ['NETO A PAGAR', '', f'${nomina.total_pagar:,.2f}']
    ]
    
    detail_table = Table(detail_data, colWidths=[3*inch, 1.5*inch, 1.5*inch])
    detail_table.setStyle(TableStyle([
        # Header
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        
        # Body
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ALIGN', (2, 1), (2, -1), 'RIGHT'),
        ('ALIGN', (1, 1), (1, -1), 'CENTER'),
        
        # Total rows
        ('BACKGROUND', (0, 7), (-1, 7), colors.HexColor('#ecf0f1')),
        ('FONTNAME', (0, 7), (-1, 7), 'Helvetica-Bold'),
        ('BACKGROUND', (0, 13), (-1, 13), colors.HexColor('#ecf0f1')),
        ('FONTNAME', (0, 13), (-1, 13), 'Helvetica-Bold'),
        
        # Final total
        ('BACKGROUND', (0, 15), (-1, 15), colors.HexColor('#27ae60')),
        ('TEXTCOLOR', (0, 15), (-1, 15), colors.whitesmoke),
        ('FONTNAME', (0, 15), (-1, 15), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 15), (-1, 15), 14),
        
        # Grid
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#bdc3c7')),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
    ]))
    
    elements.append(detail_table)
    
    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    return buffer


def generar_excel_nominas(nominas, empleados_dict):
    """
    Generate Excel report for multiple payroll records
    
    Args:
        nominas: List of Nomina objects
        empleados_dict: Dictionary mapping empleado_id to Empleado object
    
    Returns:
        BytesIO buffer with Excel content
    """
    buffer = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Nóminas"
    
    # Define styles
    header_fill = PatternFill(start_color="3498db", end_color="3498db", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    center_alignment = Alignment(horizontal="center", vertical="center")
    right_alignment = Alignment(horizontal="right", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = [
        'ID', 'Empleado', 'Periodo', 'Días Trab.', 'Hrs Extra', 'Faltas', 'Tardes',
        'Sueldo Base', 'Pago Extras', 'Ajuste Faltas', 'Bonos', 
        'Desc. Tardes', 'Desc. Préstamo', 'Caja Ahorro', 'Total Desc.', 'TOTAL A PAGAR'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = border
    
    # Data rows
    for row_idx, nomina in enumerate(nominas, 2):
        empleado = empleados_dict.get(nomina.empleado_id)
        empleado_nombre = empleado.nombre if empleado else 'N/A'
        
        data = [
            nomina.id,
            empleado_nombre,
            nomina.periodo,
            nomina.dias_trabajados,
            nomina.horas_extra,
            nomina.faltas,
            nomina.tardes,
            nomina.sueldo_base,
            nomina.pago_extras,
            nomina.ajuste_faltas,
            nomina.total_bonos,
            nomina.descuento_tardes,
            nomina.descuento_prestamo,
            nomina.caja_ahorro,
            nomina.total_descuentos,
            nomina.total_pagar
        ]
        
        for col_idx, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            
            # Apply number formatting to monetary columns
            if col_idx >= 8:
                cell.number_format = '$#,##0.00'
                cell.alignment = right_alignment
            elif col_idx >= 4 and col_idx <= 7:
                cell.alignment = center_alignment
    
    # Adjust column widths
    column_widths = [8, 20, 20, 10, 10, 8, 8, 12, 12, 12, 10, 12, 12, 12, 12, 15]
    for col_idx, width in enumerate(column_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
    
    # Save to buffer
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_resumen_periodo(nominas, periodo):
    """
    Generate summary report for a payroll period
    
    Args:
        nominas: List of Nomina objects for the period
        periodo: Period string
    
    Returns:
        Dictionary with summary statistics
    """
    if not nominas:
        return {
            'periodo': periodo,
            'total_empleados': 0,
            'total_sueldo_base': 0,
            'total_extras': 0,
            'total_bonos': 0,
            'total_descuentos': 0,
            'total_pagar': 0
        }
    
    return {
        'periodo': periodo,
        'total_empleados': len(nominas),
        'total_sueldo_base': sum(n.sueldo_base for n in nominas),
        'total_extras': sum(n.pago_extras for n in nominas),
        'total_bonos': sum(n.total_bonos for n in nominas),
        'total_descuentos': sum(n.total_descuentos for n in nominas),
        'total_pagar': sum(n.total_pagar for n in nominas)
    }
